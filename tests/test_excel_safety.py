import io

import openpyxl

import batch_ui
import resume_core as rc
import xcelgrad_sales
import xcelgrad_tech
from test_resume_core import _docx

# The skills line from a real resume: a vertical tab and other control characters sat between the skills.
DIRTY = ("Dev Sharma\ndev@example.com | +91 98765 43210\nSKILLS\n"
         "Quality Assurance (QA)\x0bQuality Control (QC)\x0cSOPs\x02, Compliance\x1b, Raw Material Testing\x00\x08 Drug Discovery\n"
         "WORK EXPERIENCE\nQA Analyst - Cipla   Jan 2020 - Present\n")


def test_control_characters_are_removed_when_reading_text():
    text = rc.extract_text("r.docx", _docx(DIRTY))
    assert not any(ord(c) < 32 and c not in "\n\t" for c in text)
    assert "Quality Assurance (QA)\nQuality Control (QC)\nSOPs, Compliance, Raw Material Testing Drug Discovery" in text


def test_dirty_resume_exports_to_excel_for_both_tools():
    for profile in (xcelgrad_tech.PROFILE, xcelgrad_sales.PROFILE):
        result = rc.process_batch([("dirty.docx", _docx(DIRTY))], profile)
        assert len(result.rows) == 1
        wb = openpyxl.load_workbook(io.BytesIO(batch_ui.generate_excel(result, "Resume_Data")))   # used to raise IllegalCharacterError
        assert wb["Resume_Data"].max_row == 2


def test_excel_writer_strips_illegal_characters_even_if_text_bypasses_the_extractor():
    result = rc.BatchResult(rows=[{"Filename": "a.pdf", "Name": "A\x0bB", "Skills": "QA\x02, QC\x1b"}],
                            failures=[("b\x03.pdf", "bad\x0b")], duplicates=[])
    wb = openpyxl.load_workbook(io.BytesIO(batch_ui.generate_excel(result, "Resume_Data")))
    assert wb["Resume_Data"]["B2"].value == "AB" and wb["Resume_Data"]["C2"].value == "QA, QC"
    assert wb["Not Processed"]["A2"].value == "b.pdf"
